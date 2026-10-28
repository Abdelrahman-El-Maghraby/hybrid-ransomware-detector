"""Builders for synthetic test samples. Nothing here is malware: every sample
only *looks like* an attack technique at the byte level (a relationship
pointing at a dead URL, a header that says MZ, random bytes, ...)."""

from __future__ import annotations

import io
import os
import struct
import zipfile
import zlib
from pathlib import Path
from typing import Dict, Optional

import docx
from docx.shared import Pt, RGBColor


def make_docx(path: Path, paragraphs=("Quarterly report", "All figures are final."),
              hidden: Optional[str] = None, white: Optional[str] = None,
              tiny: Optional[str] = None) -> Path:
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if hidden:
        d.add_paragraph().add_run(hidden).font.hidden = True
    if white:
        d.add_paragraph().add_run(white).font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if tiny:
        d.add_paragraph().add_run(tiny).font.size = Pt(1)
    d.save(str(path))
    return path


def rewrite_zip(src: Path, dst: Path, replace: Dict[str, bytes]) -> Path:
    """Copy an OOXML/ZIP file, replacing or adding the given members."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename in replace:
                continue
            zout.writestr(item, zin.read(item.filename))
        for name, data in replace.items():
            zout.writestr(name, data)
    return dst


def rels_xml(*relationships: str) -> bytes:
    body = "".join(relationships)
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f"{body}</Relationships>").encode()


def rel(rid: str, rtype: str, target: str, external: bool = True) -> str:
    mode = ' TargetMode="External"' if external else ""
    return (f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/'
            f'2006/relationships/{rtype}" Target="{target}"{mode}/>')


def make_pdf(text: str = "Hello", invisible: Optional[str] = None) -> bytes:
    """Minimal single-page PDF; `invisible` text is drawn with render mode 3."""
    ops = [f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"]
    if invisible:
        ops.append(f"BT /F1 12 Tf 3 Tr 72 700 Td ({invisible}) Tj ET")
    content = "\n".join(ops).encode("latin-1")
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content
            + b"\nendstream"]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
              .encode())
    return out.getvalue()


# ---------------------------------------------------------------- ZipCrypto
class _ZipCryptoKeys:
    """Traditional PKWARE stream cipher (what `zip -P` produces). Implemented
    here so the encrypted-archive tests run on every OS without extra tools."""

    def __init__(self, password: bytes):
        self.k = [0x12345678, 0x23456789, 0x34567890]
        for c in password:
            self.update(c)

    @staticmethod
    def _crc(crc: int, b: int) -> int:
        return zlib.crc32(bytes([b]), crc ^ 0xFFFFFFFF) ^ 0xFFFFFFFF

    def update(self, c: int) -> None:
        self.k[0] = self._crc(self.k[0], c)
        self.k[1] = (self.k[1] + (self.k[0] & 0xFF)) & 0xFFFFFFFF
        self.k[1] = (self.k[1] * 134775813 + 1) & 0xFFFFFFFF
        self.k[2] = self._crc(self.k[2], (self.k[1] >> 24) & 0xFF)

    def encrypt(self, data: bytes) -> bytes:
        out = bytearray()
        for p in data:
            t = (self.k[2] | 2) & 0xFFFF
            out.append(p ^ (((t * (t ^ 1)) >> 8) & 0xFF))
            self.update(p)
        return bytes(out)


def make_zipcrypto_zip(path: Path, name: str, data: bytes, password: str) -> Path:
    """Single STORED member encrypted with ZipCrypto."""
    crc = zlib.crc32(data) & 0xFFFFFFFF
    header = bytearray(os.urandom(11)) + bytes([(crc >> 24) & 0xFF])
    payload = _ZipCryptoKeys(password.encode()).encrypt(bytes(header) + data)
    fname = name.encode()
    dos_time, dos_date = 0, (2026 - 1980) << 9 | 1 << 5 | 1
    local = struct.pack("<IHHHHHIIIHH", 0x04034B50, 20, 0x1, 0, dos_time, dos_date, crc,
                        len(payload), len(data), len(fname), 0) + fname
    central = struct.pack("<IHHHHHHIIIHHHHHII", 0x02014B50, 20, 20, 0x1, 0, dos_time, dos_date,
                          crc, len(payload), len(data), len(fname), 0, 0, 0, 0, 0, 0) + fname
    cd_offset = len(local) + len(payload)
    eocd = struct.pack("<IHHHHIIH", 0x06054B50, 0, 0, 1, 1, len(central), cd_offset, 0)
    path.write_bytes(local + payload + central + eocd)
    return path


# --------------------------------------------------------------------- PE files
SCN_CODE_RX = 0x60000020      # code, execute, read
SCN_DATA_RW = 0xC0000040      # initialised data, read, write
SCN_RWX = 0xE0000020


def make_pe(sections=((".text", b"\x90" * 0x400, SCN_CODE_RX),), imports=None, bits=32,
            kind="exe", overlay=b"", cert=b"", strings=b"", delay_imports=None,
            delay_va_form=False) -> bytes:
    """A structurally valid PE file built from scratch (nothing runs: the code
    sections are NOPs). `imports` maps a DLL name to imported function names;
    `strings` lands in a data section; `cert` becomes the Authenticode
    directory, appended after `overlay` as signed files have it; `delay_imports`
    (same shape as `imports`) becomes the delay-load table (data directory 13)."""
    plus = bits == 64
    file_align, sect_align = 0x200, 0x1000
    secs = [list(s) for s in sections]
    if strings:
        secs.append([".rdata", strings, 0x40000040])
    n = len(secs) + (1 if imports else 0) + (1 if delay_imports else 0)
    opt_size = 240 if plus else 224
    table = 0x80 + 24 + opt_size
    headers_end = _align(table + 40 * n, file_align)

    def va(i):
        return sect_align * (i + 1)

    if imports:                                   # build .idata at its final RVA
        base = va(len(secs))
        secs.append([".idata", _idata(imports, base, plus), SCN_DATA_RW])
    didata = None
    if delay_imports:
        didata = len(secs)
        secs.append([".didat", _didata(delay_imports, va(len(secs)), plus, delay_va_form),
                     SCN_DATA_RW])
    raw, layout, off = bytearray(), [], headers_end
    for i, (name, data, chars) in enumerate(secs):
        size = _align(len(data), file_align)
        layout.append((name, va(i), len(data), off, size, chars))
        raw += data + bytes(size - len(data))
        off += size
    body_end = headers_end + len(raw)

    h = bytearray(headers_end)
    h[0:2] = b"MZ"
    struct.pack_into("<I", h, 0x3C, 0x80)
    h[0x80:0x84] = b"PE\x00\x00"
    chars = 0x0102 | (0x2000 if kind == "dll" else 0)
    struct.pack_into("<HHIIIHH", h, 0x84, 0x8664 if plus else 0x14C, len(layout), 0, 0, 0,
                     opt_size, chars)
    opt = 0x98
    struct.pack_into("<H", h, opt, 0x20B if plus else 0x10B)
    struct.pack_into("<II", h, opt + 32, sect_align, file_align)
    struct.pack_into("<H", h, opt + 68, 1 if kind == "driver" else 2)
    struct.pack_into("<I", h, opt + (108 if plus else 92), 16)
    dirs = opt + (112 if plus else 96)
    if imports:
        idata = [s for s in layout if s[0] == ".idata"][0]
        struct.pack_into("<II", h, dirs + 8, idata[1], idata[2])
    if didata is not None:
        struct.pack_into("<II", h, dirs + 13 * 8, layout[didata][1], layout[didata][2])
    if cert:
        struct.pack_into("<II", h, dirs + 32, body_end + len(overlay), len(cert))
    for i, (name, v, vsize, ptr, size, c) in enumerate(layout):
        o = table + 40 * i
        h[o:o + 8] = name.encode()[:8].ljust(8, b"\x00")
        struct.pack_into("<IIII", h, o + 8, vsize, v, size, ptr)
        struct.pack_into("<I", h, o + 36, c)
    return bytes(h) + bytes(raw) + overlay + cert


def _didata(imports, base, plus, va_form=False):
    """Delay-load descriptors (32 bytes: attributes, DLL name, module handle, IAT,
    name table, ...) followed by the same lookup tables and names as .idata."""
    width = 8 if plus else 4
    dlls = list(imports.items())
    desc_size = 32 * (len(dlls) + 1)
    thunk_len = sum((len(f) + 1) * width for _, f in dlls)
    names_start = desc_size + thunk_len
    descs, thunks_blob, names_blob = bytearray(), bytearray(), bytearray()
    for dll, funcs in dlls:
        int_rva = base + desc_size + len(thunks_blob)
        for f in funcs:
            entry_rva = base + names_start + len(names_blob)
            names_blob += b"\x00\x00" + f.encode() + b"\x00"
            if len(names_blob) % 2:
                names_blob += b"\x00"
            thunks_blob += entry_rva.to_bytes(width, "little")
        thunks_blob += bytes(width)
        dll_rva = base + names_start + len(names_blob)
        names_blob += dll.encode() + b"\x00"
        descs += struct.pack("<IIIIIIII", 0 if va_form else 1, dll_rva, 0, int_rva, int_rva,
                             0, 0, 0)
    descs += bytes(32)
    return bytes(descs + thunks_blob + names_blob)


def _align(n, a):
    return (n + a - 1) // a * a


def _idata(imports, base, plus):
    """Import descriptors, lookup tables, hint/name entries and DLL names."""
    width = 8 if plus else 4
    dlls = list(imports.items())
    desc_size = 20 * (len(dlls) + 1)
    names_blob, thunks_blob = bytearray(), bytearray()
    thunk_start = desc_size
    thunk_len = sum((len(f) + 1) * width for _, f in dlls)
    names_start = thunk_start + thunk_len
    descs = bytearray()
    for dll, funcs in dlls:
        ilt_rva = base + thunk_start + len(thunks_blob)
        for f in funcs:
            entry_rva = base + names_start + len(names_blob)
            names_blob += b"\x00\x00" + f.encode() + b"\x00"
            if len(names_blob) % 2:
                names_blob += b"\x00"
            thunks_blob += entry_rva.to_bytes(width, "little")
        thunks_blob += bytes(width)
        dll_rva = base + names_start + len(names_blob)
        names_blob += dll.encode() + b"\x00"
        descs += struct.pack("<IIIII", ilt_rva, 0, 0, dll_rva, ilt_rva)
    descs += bytes(20)
    return bytes(descs + thunks_blob + names_blob)

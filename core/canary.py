"""
Ransomware canary tripwires (honeyfiles).

WHY THIS EXISTS
---------------
Every other layer in this project looks at a file *before* it runs. Canaries
catch the case where something already got through: ransomware has to read,
encrypt and usually rename or delete user documents, and it walks folders in a
fairly predictable order. A canary is a decoy document planted where that walk
will reach it early (the `!000_` prefix sorts before ordinary names). Nobody has
a reason to touch it, so if one is deleted, rewritten, renamed or replaced with
high-entropy data, that is a high-confidence signal that encryption is in
progress - MITRE ATT&CK T1486 (Data Encrypted for Impact).

The registry (a small JSON file in the app data directory) records each
canary's path, size and SHA-256. `verify()` re-checks them; the background
scanner does this every few minutes, which turns the tool from "scan every 6
hours" into a near-real-time tripwire.

Canaries are never scanned as regular files, and `remove()` cleans them up.
If you delete one yourself, the next verify reports it - re-run deploy.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import os
import platform
import secrets
import shutil
import subprocess
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Set, Union

from detectors.file_integrity import (
    RANSOM_NOTE_RE,
    detect_format,
    looks_encrypted,
)

logger = logging.getLogger("RansomwareScanner.canary")

PathLike = Union[str, Path]

# (file name, generator kind). The "!000_" prefix makes them sort first in
# alphabetical directory walks, which is how most ransomware enumerates files.
CANARY_TEMPLATES = [
    ("!000_Finance_Records_Backup.docx", "docx"),
    ("!000_Signed_Contracts_Archive.pdf", "pdf"),
    ("!000_Asset_Inventory.csv", "csv"),
]

EXPECTED_FORMAT = {"docx": "zip", "pdf": "pdf", "csv": "text"}


def _norm(path: PathLike) -> str:
    return os.path.normcase(os.path.abspath(str(path)))


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


# --------------------------------------------------------------- generators
def _make_docx(token: str) -> bytes:
    import docx  # python-docx is a core dependency of the scanner
    d = docx.Document()
    d.add_heading("Finance Records - Backup Index", level=1)
    d.add_paragraph("Consolidated ledger exports, vendor payment batches and "
                    "reconciliation notes. Do not move this file; it is referenced "
                    "by the quarterly close checklist.")
    for i in range(1, 6):
        d.add_paragraph(f"Batch {i:02d}: archived {time.strftime('%Y-%m')} - ref {token[:8]}-{i}")
    d.add_paragraph(f"Index id: {token}")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _make_pdf(token: str) -> bytes:
    """A tiny but valid single-page PDF with real text (built by hand)."""
    lines = ["Signed Contracts - Archive Index",
             f"Document set {token[:8]}",
             "Retention: 7 years. Contact legal before modifying."]
    text_ops = ["BT", "/F1 12 Tf", "72 760 Td", "14 TL"]
    for line in lines:
        safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        text_ops.append(f"({safe}) Tj T*")
    text_ops.append("ET")
    content = "\n".join(text_ops).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Producer (" + token.encode() + b") >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for num, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{num} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R /Info 6 0 R >>\n"
              f"startxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def _make_csv(token: str) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["asset_tag", "type", "owner", "location", "purchased"])
    for i in range(1, 25):
        w.writerow([f"AT-{token[:4].upper()}-{i:04d}", ["Laptop", "Monitor", "Phone"][i % 3],
                    f"dept-{i % 7}", f"Floor {1 + i % 4}", f"20{18 + i % 7}-0{1 + i % 9}-15"])
    return buf.getvalue().encode("utf-8")


_GENERATORS = {"docx": _make_docx, "pdf": _make_pdf, "csv": _make_csv}


# ------------------------------------------------------------------ results
@dataclass
class CanaryFinding:
    path: str
    status: str            # deleted | modified | encrypted | renamed | ransom_note
    evidence: str

    def as_dict(self) -> Dict[str, str]:
        return {"path": self.path, "status": self.status, "evidence": self.evidence}


@dataclass
class CanaryReport:
    deployed: int = 0
    intact: int = 0
    findings: List[CanaryFinding] = field(default_factory=list)
    checked_at: str = ""

    @property
    def triggered(self) -> bool:
        return bool(self.findings)

    def summary(self) -> str:
        if not self.deployed:
            return "No canaries deployed."
        if not self.triggered:
            return f"{self.intact}/{self.deployed} canaries intact."
        kinds = ", ".join(sorted({f.status for f in self.findings}))
        return (f"TRIPWIRE TRIGGERED - {len(self.findings)} finding(s) ({kinds}); "
                f"{self.intact}/{self.deployed} canaries intact.")

    def as_dict(self) -> Dict[str, Any]:
        return {"deployed": self.deployed, "intact": self.intact,
                "triggered": self.triggered, "checked_at": self.checked_at,
                "summary": self.summary(),
                "findings": [f.as_dict() for f in self.findings]}


# ------------------------------------------------------------------ manager
class CanaryManager:
    """Deploys, verifies and removes canary files; state lives in a JSON registry."""

    def __init__(self, registry_path: PathLike):
        self.registry_path = Path(registry_path)

    # -- registry ------------------------------------------------------------
    def _load(self) -> List[Dict[str, Any]]:
        try:
            data = json.loads(self.registry_path.read_text(encoding="utf-8"))
            entries = data.get("canaries", []) if isinstance(data, dict) else []
            return [e for e in entries if isinstance(e, dict) and e.get("path")]
        except FileNotFoundError:
            return []
        except Exception as e:
            logger.error("Canary registry unreadable (%s): %s", self.registry_path, e)
            return []

    def _save(self, entries: List[Dict[str, Any]]) -> None:
        self.registry_path.parent.mkdir(parents=True, exist_ok=True)
        tmp = self.registry_path.with_suffix(".tmp")
        tmp.write_text(json.dumps({"version": 1, "canaries": entries}, indent=2),
                       encoding="utf-8")
        os.replace(tmp, self.registry_path)

    def entries(self) -> List[Dict[str, Any]]:
        return self._load()

    def paths(self) -> Set[str]:
        """Normalised paths of every registered canary (for scan exclusion)."""
        return {_norm(e["path"]) for e in self._load()}

    # -- deploy ----------------------------------------------------------------
    def deploy(self, directories: Iterable[PathLike]) -> List[Dict[str, Any]]:
        """Plant the canary set in each existing directory. Never overwrites an
        existing file; already-registered canaries are left as they are."""
        entries = self._load()
        registered = {_norm(e["path"]) for e in entries}
        created: List[Dict[str, Any]] = []
        for directory in directories:
            d = Path(directory).expanduser()
            if not d.is_dir():
                logger.warning("Canary deploy skipped - not a directory: %s", d)
                continue
            for name, kind in CANARY_TEMPLATES:
                target = d / name
                if _norm(target) in registered:
                    continue
                if target.exists():
                    logger.warning("Canary deploy skipped - a non-canary file exists at %s", target)
                    continue
                try:
                    data = _GENERATORS[kind](secrets.token_hex(16))
                    with open(target, "xb") as fh:      # 'x' = never clobber
                        fh.write(data)
                except Exception as e:
                    logger.error("Could not create canary %s: %s", target, e)
                    continue
                entry = {"path": str(target.resolve()), "kind": kind,
                         "sha256": hashlib.sha256(data).hexdigest(), "size": len(data),
                         "deployed_at": time.strftime("%Y-%m-%d %H:%M:%S")}
                entries.append(entry)
                registered.add(_norm(target))
                created.append(entry)
        self._save(entries)
        if created:
            logger.info("Deployed %d canary file(s).", len(created))
        return created

    # -- verify ----------------------------------------------------------------
    def verify(self) -> CanaryReport:
        entries = self._load()
        report = CanaryReport(deployed=len(entries),
                              checked_at=time.strftime("%Y-%m-%d %H:%M:%S"))
        note_dirs: Set[str] = set()
        for e in entries:
            path = Path(e["path"])
            note_dirs.add(str(path.parent))
            if not path.exists():
                sibling = self._find_renamed(path)
                if sibling is not None:
                    report.findings.append(CanaryFinding(
                        str(path), "renamed",
                        f"Canary disappeared and '{sibling.name}' appeared beside it - "
                        f"typical of ransomware appending its extension."))
                else:
                    report.findings.append(CanaryFinding(
                        str(path), "deleted", "Canary file was deleted."))
                continue
            try:
                digest = _sha256(path)
            except OSError as ex:
                report.findings.append(CanaryFinding(str(path), "modified",
                                                     f"Canary unreadable: {ex}"))
                continue
            if digest == e.get("sha256"):
                report.intact += 1
                continue
            try:
                with path.open("rb") as fh:
                    head = fh.read(65536)
            except OSError:
                head = b""
            expected = EXPECTED_FORMAT.get(e.get("kind", ""), "")
            fmt = detect_format(head)
            # Random-looking bytes, or a document that turned into an
            # unrecognisable binary blob, means it was most likely encrypted.
            if looks_encrypted(head) or (expected and fmt == "unknown"):
                report.findings.append(CanaryFinding(
                    str(path), "encrypted",
                    f"Canary content was rewritten with high-entropy data "
                    f"(format now '{fmt}', expected '{expected}')."))
            else:
                report.findings.append(CanaryFinding(
                    str(path), "modified", "Canary content changed (SHA-256 mismatch)."))
        # Ransom notes dropped next to canaries are a second, independent signal.
        for d in sorted(note_dirs):
            try:
                for child in Path(d).iterdir():
                    if child.is_file() and RANSOM_NOTE_RE.search(child.name):
                        report.findings.append(CanaryFinding(
                            str(child), "ransom_note",
                            "A file named like a ransom note appeared in a canary folder."))
            except OSError:
                continue
        if report.triggered:
            logger.critical("CANARY %s", report.summary())
            for f in report.findings:
                logger.critical("  [%s] %s - %s", f.status, f.path, f.evidence)
        return report

    @staticmethod
    def _find_renamed(path: Path) -> Optional[Path]:
        try:
            for child in path.parent.iterdir():
                if child.name == path.name:
                    continue
                appended = child.name.startswith(path.name + ".")          # x.docx.lockbit
                swapped = (child.name.startswith(path.stem + ".")           # x.locked
                           and child.suffix.lower() != path.suffix.lower())
                if appended or swapped:
                    return child
        except OSError:
            return None
        return None

    # -- remove ----------------------------------------------------------------
    def remove(self) -> int:
        """Delete every registered canary file and clear the registry."""
        entries = self._load()
        removed = 0
        for e in entries:
            try:
                Path(e["path"]).unlink()
                removed += 1
            except FileNotFoundError:
                pass
            except OSError as ex:
                logger.error("Could not remove canary %s: %s", e["path"], ex)
        self._save([])
        return removed


def desktop_notify(title: str, message: str) -> bool:
    """Best-effort desktop notification (macOS / Linux). Never raises."""
    system = platform.system()
    try:
        if system == "Darwin" and shutil.which("osascript"):
            def esc(s: str) -> str:
                return s.replace("\\", "\\\\").replace('"', '\\"')
            script = f'display notification "{esc(message)}" with title "{esc(title)}"'
            subprocess.run(["osascript", "-e", script], capture_output=True, timeout=10)
            return True
        if system == "Linux" and shutil.which("notify-send"):
            subprocess.run(["notify-send", "--urgency=critical", title, message],
                           capture_output=True, timeout=10)
            return True
    except Exception as e:  # pragma: no cover
        logger.debug("Desktop notification failed: %s", e)
    return False
